"""
Forecast-run QA workbook for planners and engine owners.

    python -m reports_tool.forecast_qa_report     -> reports/forecast_qa_report.xlsx
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.gam_engine import run
from engine.seed import build

ROOT = Path(__file__).resolve().parent.parent
HEAD = PatternFill("solid", fgColor="1F4E79")
OK, BAD = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
WAPE_TARGET = 15.0


def sheet(wb, title, header, rows, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append(list(r))
        if status_col is not None:
            ws.cell(ws.max_row, status_col).fill = OK if r[status_col - 1] == "PASS" else BAD
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(len(str(c.value or "")) for c in col) + 2, 60)
    return ws


def main(out: Path = ROOT / "reports" / "forecast_qa_report.xlsx") -> Path:
    db = build(ROOT / "db" / "forecast.db")
    rid = run(db)
    con = sqlite3.connect(db)
    wb = Workbook()
    wb.remove(wb.active)

    checks = []
    for f in sorted((ROOT / "sql" / "validation").glob("*.sql")):
        n = len(con.execute(f.read_text()).fetchall())
        checks.append((f.stem, n, "PASS" if n == 0 else "FAIL"))
    run_row = con.execute("SELECT run_id, model_version, history_end, horizon_weeks, status FROM engine_run WHERE run_id=?", (rid,)).fetchone()
    sheet(wb, "Run summary", ["Run", "Model version", "History end", "Horizon (weeks)", "Status"], [run_row])
    sheet(wb, "Validation checks", ["Check", "Offending rows", "Status"], checks, status_col=3)
    acc = con.execute("""SELECT product_id, location_id, mape, wape, bias,
                                CASE WHEN wape <= ? THEN 'PASS' ELSE 'REVIEW' END FROM forecast_accuracy
                         WHERE run_id = ? ORDER BY wape DESC""", (WAPE_TARGET, rid)).fetchall()
    ws = sheet(wb, "Accuracy", ["Product", "Location", "MAPE %", "WAPE %", "Bias %", f"WAPE <= {WAPE_TARGET}%"], acc)
    for i, r in enumerate(acc, start=2):
        ws.cell(i, 6).fill = OK if r[5] == "PASS" else BAD
    sheet(wb, "Exceptions", ["Product", "Location", "Reason", "Detail"],
          con.execute("SELECT product_id, location_id, reason_code, detail FROM engine_exception WHERE run_id=?", (rid,)).fetchall())
    sheet(wb, "Forecast by location", ["Location", "Week", "Total forecast"],
          con.execute("SELECT location_id, week_start, total_forecast FROM v_forecast_by_location WHERE run_id=? ORDER BY 1, 2", (rid,)).fetchall())
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print("Report written to", main())
